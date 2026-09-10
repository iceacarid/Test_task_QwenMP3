import json
import sys
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


def add_approval_block(doc):
    """Блок УТВЕРЖДАЮ — административные поля не извлекаются из аудио,
    оставлены плейсхолдерами в квадратных скобках, как в шаблоне."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run("УТВЕРЖДАЮ").bold = True

    for text in ("[Должность]", "[Наименование организации]", "______________ [Ф.И.О.]", "«____» ____________ 20__ г."):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.add_run(text)


def add_title(doc):
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("П Р О Т О К О Л")
    run.bold = True
    run.font.size = Pt(16)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("совещания при [должность / Ф.И.О.]")

    number_line = doc.add_paragraph()
    number_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    number_line.add_run("№___от __.__.20__")

    city_line = doc.add_paragraph()
    city_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    city_line.add_run("[город / населённый пункт]")


def add_closing_block(doc):
    """Подвал "Протокол вёл" — как и УТВЕРЖДАЮ, поле не определяется из аудио."""
    doc.add_paragraph()
    doc.add_paragraph().add_run("Протокол вёл:").bold = True
    doc.add_paragraph("[Должность]")
    doc.add_paragraph("______________")
    doc.add_paragraph("[Ф.И.О.]")


def add_meta(doc, meta):
    doc.add_paragraph()
    time = meta.get("time") or "не указано в записи"
    place = meta.get("place") or "не указано в записи"
    participants = meta.get("participants") or []
    participants_str = ", ".join(participants) if participants else "не указаны в записи"

    p = doc.add_paragraph()
    p.add_run("Время проведения: ").bold = True
    p.add_run(time)

    p = doc.add_paragraph()
    p.add_run("Место проведения: ").bold = True
    p.add_run(place)

    p = doc.add_paragraph()
    p.add_run("Присутствуют: ").bold = True
    p.add_run(participants_str)


def add_agenda_items(doc, agenda_items):
    heading = doc.add_paragraph()
    heading.add_run("Заслушали:").bold = True

    for item in agenda_items:
        doc.add_paragraph()
        speaker = item.get("speaker") or "докладчик не указан"

        speaker_p = doc.add_paragraph()
        speaker_p.add_run(speaker).bold = True
        timestamp = item.get("audio_timestamp")
        if timestamp:
            ts_run = speaker_p.add_run(f"  [{timestamp}]")
            ts_run.italic = True

        doc.add_paragraph(item.get("topic_summary") or "")

        decisions = item.get("decisions") or []
        if decisions:
            p = doc.add_paragraph()
            p.add_run("Решение:").bold = True
            for d in decisions:
                line = f"– {d.get('text') or ''}"
                extra = []
                if d.get("responsible"):
                    extra.append(f"ответственный: {d['responsible']}")
                if d.get("deadline"):
                    extra.append(f"срок: {d['deadline']}")
                if extra:
                    line += " (" + ", ".join(extra) + ")"
                doc.add_paragraph(line)
        else:
            p = doc.add_paragraph()
            run = p.add_run("Решение не принято.")
            run.italic = True

        needs_review = item.get("needs_review") or []
        if needs_review:
            p = doc.add_paragraph()
            run = p.add_run("Требует проверки:")
            run.italic = True
            run.bold = True
            for fragment in needs_review:
                p2 = doc.add_paragraph()
                r2 = p2.add_run(f"— {fragment}")
                r2.italic = True


def build_protocol_docx(protocol: dict, output_path: str):
    """Собирает DOCX-протокол из словаря (результат generate_protocol) и сохраняет по указанному пути."""
    doc = Document()
    add_approval_block(doc)
    add_title(doc)
    add_meta(doc, protocol.get("meta") or {})
    add_agenda_items(doc, protocol.get("agenda_items") or [])
    add_closing_block(doc)
    doc.save(output_path)


if __name__ == "__main__":
    input_path = sys.argv[1] if len(sys.argv) > 1 else "protocol.json"
    output_path = sys.argv[2] if len(sys.argv) > 2 else "protocol.docx"
    with open(input_path, "r", encoding="utf-8") as f:
        protocol_data = json.load(f)
    build_protocol_docx(protocol_data, output_path)
    print(f"Готово: {output_path}")
